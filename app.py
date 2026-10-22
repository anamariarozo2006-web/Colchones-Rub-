"""Colchones Rubí - Control de inventario (Flask + SQLite).
Instalar:   pip install -r requirements.txt
Ejecutar:   python app.py            -> http://127.0.0.1:5000
Producción: waitress-serve --listen=0.0.0.0:8000 app:app   (detrás de HTTPS; defina COOKIE_SECURE=1)
Usuarios:   python app.py usuario <documento> "<Nombre>" [admin|operario]   (crea o cambia contraseña)
"""
import os, sys, time, sqlite3, secrets
from datetime import datetime
from functools import wraps
from flask import Flask, g, request, session, redirect, url_for, jsonify, render_template, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
import io

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get("DB_PATH", os.path.join(BASE, "inventario.db"))
KEYFILE = os.path.join(BASE, ".secret_key")
if not os.path.exists(KEYFILE):
    open(KEYFILE, "w").write(secrets.token_hex(32))

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or open(KEYFILE).read()
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE") == "1",
                  PERMANENT_SESSION_LIFETIME=8 * 3600)

SCHEMA = """
CREATE TABLE IF NOT EXISTS usuarios(id INTEGER PRIMARY KEY, documento TEXT UNIQUE NOT NULL, nombre TEXT NOT NULL,
  rol TEXT NOT NULL DEFAULT 'operario', pw_hash TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS productos(codigo TEXT PRIMARY KEY, nombre TEXT NOT NULL, categoria TEXT, costo REAL NOT NULL,
  detal REAL DEFAULT 0, mayor REAL DEFAULT 0, stock INTEGER NOT NULL DEFAULT 0 CHECK(stock>=0), minimo INTEGER DEFAULT 5);
CREATE TABLE IF NOT EXISTS movimientos(id INTEGER PRIMARY KEY, fecha TEXT, usuario TEXT, tipo TEXT, codigo TEXT,
  producto TEXT, cantidad INTEGER, valor_unit REAL, valor_total REAL, costo_total REAL);
"""
SEED = [
("COL-001","Colchón Sencillo Espuma","Colchones",180000,290000,260000,24),
("COL-002","Colchón Semidoble Espuma","Colchones",230000,365000,330000,18),
("COL-003","Colchón Doble Resortes","Colchones",420000,690000,620000,12),
("COL-004","Colchón Queen Ortopédico","Colchones",560000,920000,830000,9),
("COL-005","Colchón King Pillow Top","Colchones",780000,1290000,1160000,6),
("BAS-001","Base Cama Semidoble","Bases",150000,245000,220000,14),
("BAS-002","Base Cama Doble","Bases",190000,310000,280000,10),
("ALM-001","Almohada Viscoelástica","Accesorios",28000,55000,46000,60),
("ALM-002","Almohada Fibra Hueca","Accesorios",12000,26000,21000,85),
("PRO-001","Protector Impermeable Doble","Accesorios",35000,69000,58000,40),
("MAT-001","Rollo Espuma D-26 (metro)","Materia prima",22000,0,0,120),
("MAT-002","Resortes Bonnell (caja)","Materia prima",95000,0,0,30),
("MAT-003","Tela Damasco (rollo)","Materia prima",210000,0,0,3)]

def init():
    c = sqlite3.connect(DB); c.executescript(SCHEMA)
    if not c.execute("SELECT 1 FROM productos").fetchone():
        c.executemany("INSERT INTO productos(codigo,nombre,categoria,costo,detal,mayor,stock) VALUES(?,?,?,?,?,?,?)", SEED)
    if not c.execute("SELECT 1 FROM usuarios").fetchone():
        pw = os.environ.get("ADMIN_PASSWORD") or secrets.token_urlsafe(9)
        c.execute("INSERT INTO usuarios(documento,nombre,rol,pw_hash) VALUES(?,?,?,?)",
                  ("1234567890", "Administrador", "admin", generate_password_hash(pw)))
        print(f"\n>>> Usuario inicial -> documento: 1234567890 | contraseña: {pw}  (cámbiela con el comando 'usuario')\n")
    c.commit(); c.close()
init()

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB, timeout=10); g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def cerrar(_):
    d = g.pop("db", None)
    if d: d.close()

def login_required(f):
    @wraps(f)
    def w(*a, **k):
        api = request.path.startswith("/api")
        if "uid" not in session:
            return (jsonify(error="Sesión expirada"), 401) if api else redirect(url_for("login"))
        if api and request.method == "POST" and not request.is_json:   # protección CSRF
            return jsonify(error="Solicitud inválida"), 400
        return f(*a, **k)
    return w

ahora = lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S")
FAIL = {}

@app.route("/login", methods=["GET", "POST"])
def login():
    err = ""
    if request.method == "POST":
        ip = request.remote_addr; n, t = FAIL.get(ip, (0, 0))
        if n >= 5 and time.time() - t < 300:
            err = "Demasiados intentos. Espere 5 minutos."
        else:
            u = db().execute("SELECT * FROM usuarios WHERE documento=?", (request.form.get("documento", "").strip(),)).fetchone()
            if u and check_password_hash(u["pw_hash"], request.form.get("password", "")):
                FAIL.pop(ip, None); session.clear()
                session.update(uid=u["id"], nombre=u["nombre"], rol=u["rol"]); session.permanent = True
                return redirect(url_for("inicio"))
            FAIL[ip] = ((n + 1) if time.time() - t < 300 else 1, time.time())
            err = "Documento o contraseña incorrectos."
    return render_template("login.html", err=err)

@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("login"))

@app.route("/")
@login_required
def inicio():
    return render_template("index.html", nombre=session["nombre"], rol=session["rol"])

VENTAS = "('VENTA DETAL','VENTA MAYOR')"
def resumen():
    r = db().execute(f"""SELECT
      COALESCE(SUM(CASE WHEN tipo='ENTRADA' THEN valor_total END),0) gasto,
      COALESCE(SUM(CASE WHEN tipo IN {VENTAS} THEN valor_total END),0) ventas,
      COALESCE(SUM(CASE WHEN tipo IN {VENTAS} THEN costo_total END),0) cogs FROM movimientos""").fetchone()
    return dict(r)

@app.route("/api/estado")
@login_required
def estado():
    ps = [dict(code=r["codigo"], name=r["nombre"], cat=r["categoria"], cost=r["costo"], detal=r["detal"],
               mayor=r["mayor"], stock=r["stock"], min=r["minimo"])
          for r in db().execute("SELECT * FROM productos ORDER BY codigo")]
    return jsonify(productos=ps, resumen=resumen())

TIPOS = {"entrada": "ENTRADA", "detal": "VENTA DETAL", "mayor": "VENTA MAYOR", "consumo": "CONSUMO"}

@app.route("/api/movimiento", methods=["POST"])
@login_required
def movimiento():
    d = request.get_json(silent=True) or {}
    tipo = d.get("tipo")
    try: n = int(d.get("cantidad"))
    except (TypeError, ValueError): n = 0
    if tipo not in TIPOS or not 0 < n <= 100000:
        return jsonify(error="Datos inválidos"), 400
    c = db(); c.execute("BEGIN IMMEDIATE")
    p = c.execute("SELECT * FROM productos WHERE codigo=?", (d.get("codigo"),)).fetchone()
    err = None
    if not p: err = "Producto no existe"
    elif tipo in ("detal", "mayor") and not p[tipo]: err = "Este producto no se vende"
    elif tipo != "entrada" and n > p["stock"]: err = f"Solo hay {p['stock']} en stock"
    if err:
        c.rollback(); return jsonify(error=err), 400
    delta = n if tipo == "entrada" else -n
    unit = p["costo"] if tipo in ("entrada", "consumo") else p[tipo]
    c.execute("UPDATE productos SET stock=stock+? WHERE codigo=?", (delta, p["codigo"]))
    c.execute("INSERT INTO movimientos(fecha,usuario,tipo,codigo,producto,cantidad,valor_unit,valor_total,costo_total) VALUES(?,?,?,?,?,?,?,?,?)",
              (ahora(), session["nombre"], TIPOS[tipo], p["codigo"], p["nombre"], n, unit, unit * n, p["costo"] * n))
    c.commit()
    return jsonify(ok=True, agotado=p["stock"] + delta == 0)

@app.route("/api/pedido", methods=["POST"])
@login_required
def pedido():
    c = db()
    for cod in (request.get_json(silent=True) or {}).get("codigos", []):
        p = c.execute("SELECT * FROM productos WHERE codigo=?", (cod,)).fetchone()
        if p:
            c.execute("INSERT INTO movimientos(fecha,usuario,tipo,codigo,producto,cantidad,valor_unit,valor_total,costo_total) VALUES(?,?,?,?,?,?,?,?,?)",
                      (ahora(), session["nombre"], "PEDIDO SOLICITADO", cod, p["nombre"], 0, p["costo"], 0, 0))
    c.commit()
    return jsonify(ok=True)

COP = '"$"#,##0'
def tabla(ws, nombre, cab, filas, anchos, fmt=()):
    ws.append(cab)
    for f in filas: ws.append(f)
    t = Table(displayName=nombre, ref=f"A1:{chr(64 + len(cab))}{len(filas) + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium3", showRowStripes=True); ws.add_table(t)
    for i, w in enumerate(anchos, 1): ws.column_dimensions[chr(64 + i)].width = w
    for col in fmt:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row: cell.number_format = COP

@app.route("/exportar")
@login_required
def exportar():
    r = resumen(); c = db()
    prods = c.execute("SELECT * FROM productos ORDER BY codigo").fetchall()
    movs = c.execute("SELECT * FROM movimientos ORDER BY id").fetchall()
    net = r["ventas"] - r["gasto"]
    wb = Workbook(); ws = wb.active; ws.title = "Resumen"; ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:B1"); ws["A1"] = "Colchones Rubí · Reporte de inventario"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor="B3131B")
    ws.row_dimensions[1].height = 30
    filas = [("Generado", ahora()), ("Generado por", session["nombre"]),
             ("Total gastado en compras", r["gasto"]), ("Total vendido", r["ventas"]),
             ("Costo de lo vendido", r["cogs"]), ("Utilidad bruta (ventas − costo vendido)", r["ventas"] - r["cogs"]),
             ("Utilidad / pérdida (ventas − gastos)", net),
             ("Valor actual del inventario", sum(p["stock"] * p["costo"] for p in prods)),
             ("Resultado", "UTILIDAD" if net >= 0 else "PÉRDIDA")]
    for i, (a, b) in enumerate(filas, 2):
        ws.cell(i, 1, a).font = Font(bold=True); x = ws.cell(i, 2, b)
        if isinstance(b, (int, float)): x.number_format = COP
    ws["B8"].font = Font(bold=True, color="1F7A45" if net >= 0 else "B3131B")
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 24
    tabla(wb.create_sheet("Inventario"), "TablaInventario",
          ["Código", "Producto", "Categoría", "Cantidad", "Costo unitario", "Precio detal", "Precio mayorista", "Valor en inventario", "Estado"],
          [[p["codigo"], p["nombre"], p["categoria"], p["stock"], p["costo"], p["detal"] or "N/A", p["mayor"] or "N/A",
            p["stock"] * p["costo"], "AGOTADO - PEDIR" if p["stock"] == 0 else "Stock bajo" if p["stock"] <= p["minimo"] else "Disponible"]
           for p in prods], [12, 34, 16, 11, 16, 15, 17, 20, 18], (5, 6, 7, 8))
    tabla(wb.create_sheet("Movimientos"), "TablaMovimientos",
          ["Fecha", "Usuario", "Tipo", "Código", "Producto", "Cantidad", "Valor unitario", "Valor total"],
          [[m["fecha"], m["usuario"], m["tipo"], m["codigo"], m["producto"], m["cantidad"], m["valor_unit"], m["valor_total"]] for m in movs]
          or [["—", "", "Sin movimientos", "", "", 0, 0, 0]], [20, 20, 22, 12, 34, 10, 16, 16], (7, 8))
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Inventario_Colchones_{datetime.now():%Y-%m-%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == "__main__":
    if len(sys.argv) >= 4 and sys.argv[1] == "usuario":
        import getpass
        pw = getpass.getpass("Contraseña (mínimo 8 caracteres): ")
        if len(pw) < 8: sys.exit("Contraseña muy corta.")
        c = sqlite3.connect(DB)
        c.execute("""INSERT INTO usuarios(documento,nombre,rol,pw_hash) VALUES(?,?,?,?) ON CONFLICT(documento)
                     DO UPDATE SET nombre=excluded.nombre, rol=excluded.rol, pw_hash=excluded.pw_hash""",
                  (sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else "operario", generate_password_hash(pw)))
        c.commit(); print("Usuario guardado.")
    else:
        app.run(debug=os.environ.get("DEBUG") == "1")
